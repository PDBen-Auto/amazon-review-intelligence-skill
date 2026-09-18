#!/usr/bin/env python3
import argparse
import csv
import html
import json
import re
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlsplit

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from PIL import Image as PILImage
except ImportError as exc:
    raise SystemExit("Install openpyxl and Pillow before building Excel: pip install openpyxl pillow") from exc


ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")
SOURCE_PATTERNS = [
    ("*_merged.json", "merged"),
    ("*_woot_max.json", "interface"),
    ("*_woot_full.json", "interface"),
    ("*_woot_basic.json", "interface"),
    ("*_browser_pages_reviews.json", "browser"),
    ("*_browser_recent_reviews.json", "browser"),
    ("*_browser_reviews.json", "browser"),
]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def safe_text(value) -> str:
    value = html.unescape(str(value or ""))
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def normalize(value) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", html.unescape(str(value or "")).lower())


def parse_rating(value) -> int:
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"([1-5])(?:\.0)?", str(value or ""))
    return int(match.group(1)) if match else 0


def truth_text(value):
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return safe_text(value)


def media_urls(value) -> list[str]:
    if not value:
        return []
    if isinstance(value, str):
        return [line.strip() for line in value.splitlines() if line.strip()]
    result = []
    for item in value:
        if isinstance(item, str):
            url = item
        elif isinstance(item, dict):
            url = item.get("src") or item.get("url") or item.get("href") or item.get("poster") or ""
        else:
            url = ""
        if url and url not in result:
            result.append(url)
    return result


def normalize_review(raw: dict, source: str) -> dict:
    images = media_urls(raw.get("ImageUrls") or raw.get("mediaImages") or raw.get("image_urls") or raw.get("图片URL"))
    videos = media_urls(raw.get("MediaUrls") or raw.get("videoElements") or raw.get("media_urls") or raw.get("视频URL"))
    return {
        "id": safe_text(raw.get("ReviewId") or raw.get("Id") or raw.get("id") or raw.get("review_id") or raw.get("评论ID")),
        "rating": parse_rating(raw.get("OverallRating") or raw.get("rating") or raw.get("评分")),
        "title": safe_text(raw.get("Title") or raw.get("title") or raw.get("标题")),
        "body": safe_text(raw.get("Text") or raw.get("body") or raw.get("text") or raw.get("评论内容") or raw.get("评论")),
        "author": safe_text(raw.get("Author") or raw.get("profile") or raw.get("author") or raw.get("评论人")),
        "date": safe_text(raw.get("OriginDescription") or raw.get("date") or raw.get("reviewDate") or raw.get("评论日期")),
        "verified": truth_text(raw.get("IsVerifiedPurchase") if "IsVerifiedPurchase" in raw else raw.get("verified") or raw.get("verified_purchase") or raw.get("是否认证购买")),
        "vine": truth_text(raw.get("IsVineReview") if "IsVineReview" in raw else raw.get("vine_review") or raw.get("Vine")),
        "helpful": raw.get("HelpfulVotes") or raw.get("helpful") or raw.get("helpful_votes") or raw.get("Helpful") or "",
        "image_urls": images,
        "video_urls": videos,
        "page_url": safe_text(raw.get("pageUrl") or raw.get("reviewUrl") or raw.get("评论链接")),
        "source": source,
        "image_paths": [],
    }


def review_key(review: dict) -> str:
    if review["id"]:
        return "id:" + review["id"]
    date_or_author = normalize(review["date"]) or normalize(review["author"])
    return "content:" + "|".join([
        normalize(review["title"]),
        normalize(review["body"][:120]),
        date_or_author,
    ])


def merge_record(existing: dict, incoming: dict) -> None:
    for field in ["id", "rating", "title", "body", "author", "date", "verified", "vine", "helpful", "page_url"]:
        if not existing.get(field) and incoming.get(field):
            existing[field] = incoming[field]
    for field in ["image_urls", "video_urls"]:
        existing[field] = list(dict.fromkeys((existing.get(field) or []) + (incoming.get(field) or [])))
    if incoming["source"] not in existing["source"].split("+"):
        existing["source"] += "+" + incoming["source"]


def extract_reviews(data):
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        return data.get("reviews") or data.get("Reviews") or []
    return []


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def collect_asin_reviews(asin_dir: Path) -> list[dict]:
    files = []
    for pattern, source in SOURCE_PATTERNS:
        files.extend((path, source) for path in sorted(asin_dir.glob(pattern)))
    if not files:
        files.extend((path, "archive") for path in sorted(asin_dir.glob("*_reviews_raw.json")))

    seen = {}
    rows = []
    for path, source in files:
        try:
            reviews = extract_reviews(read_json(path))
        except Exception:
            continue
        for raw in reviews:
            review = normalize_review(raw, source)
            if not review["body"] and not review["title"]:
                continue
            key = review_key(review)
            if key in seen:
                merge_record(seen[key], review)
            else:
                seen[key] = review
                rows.append(review)
    return rows


def image_ext(url: str) -> str:
    ext = Path(urlsplit(url).path).suffix.lower()
    return ext if ext in {".jpg", ".jpeg", ".png", ".webp"} else ".jpg"


def download_image(url: str, target: Path) -> str:
    if target.exists() and target.stat().st_size > 0:
        return str(target)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.amazon.com/"})
        with urllib.request.urlopen(req, timeout=10) as response:
            target.write_bytes(response.read())
        return str(target) if target.stat().st_size > 0 else ""
    except Exception:
        return ""


def make_excel_image(path: str, max_w=118, max_h=88):
    if not path or not Path(path).exists():
        return None
    try:
        with PILImage.open(path) as image:
            width, height = image.size
        xl = XLImage(path)
        ratio = min(max_w / max(width, 1), max_h / max(height, 1), 1)
        xl.width = int(width * ratio)
        xl.height = int(height * ratio)
        return xl
    except Exception:
        return None


def add_header_style(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_asin_csv(asin_dir: Path, asin: str, rows: list[dict]) -> None:
    headers = ["序号", "ASIN", "评论ID", "评分", "标题", "评论内容", "评论人", "评论日期", "是否认证购买", "Vine", "Helpful", "图片数量", "图片URL", "视频URL", "来源"]
    with (asin_dir / f"{asin}_reviews.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for index, row in enumerate(rows, 1):
            writer.writerow({
                "序号": index, "ASIN": asin, "评论ID": row["id"], "评分": row["rating"], "标题": row["title"],
                "评论内容": row["body"], "评论人": row["author"], "评论日期": row["date"],
                "是否认证购买": row["verified"], "Vine": row["vine"], "Helpful": row["helpful"],
                "图片数量": len(row["image_urls"]), "图片URL": "\n".join(row["image_urls"]),
                "视频URL": "\n".join(row["video_urls"]), "来源": row["source"],
            })


def main() -> None:
    parser = argparse.ArgumentParser(description="Build one deduplicated Amazon review Excel workbook")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output")
    parser.add_argument("--download-images", action="store_true")
    parser.add_argument("--max-images-per-review", type=int, default=3)
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    input_dir = Path(args.input_dir).resolve()
    output = Path(args.output).resolve() if args.output else input_dir / "ASIN_reviews_ALL.xlsx"
    asin_dirs = sorted(path for path in input_dir.iterdir() if path.is_dir() and ASIN_RE.fullmatch(path.name.upper()))
    all_rows = []
    grouped = {}
    for asin_dir in asin_dirs:
        asin = asin_dir.name.upper()
        rows = collect_asin_reviews(asin_dir)
        grouped[asin] = rows
        write_asin_csv(asin_dir, asin, rows)
        for index, row in enumerate(rows, 1):
            row.update({"asin": asin, "asin_index": index, "global_index": len(all_rows) + 1})
            all_rows.append(row)

    if args.download_images:
        jobs = []
        for row in all_rows:
            for ordinal, url in enumerate(row["image_urls"][:max(args.max_images_per_review, 0)], 1):
                target = input_dir / row["asin"] / "media" / "images" / f"{row['asin']}_{row['asin_index']:04d}_{ordinal}{image_ext(url)}"
                jobs.append((row, ordinal, url, target))
        with ThreadPoolExecutor(max_workers=max(args.workers, 1)) as pool:
            futures = {pool.submit(download_image, url, target): (row, ordinal) for row, ordinal, url, target in jobs}
            for future in as_completed(futures):
                row, ordinal = futures[future]
                path = future.result()
                if path:
                    row["image_paths"].append((ordinal, path))
        for row in all_rows:
            row["image_paths"] = [path for _, path in sorted(row["image_paths"])]

    summary = []
    for asin_dir in asin_dirs:
        asin = asin_dir.name.upper()
        rows = grouped[asin]
        stars = Counter(row["rating"] for row in rows)
        summary.append({
            "ASIN": asin, "评论数": len(rows), "5星": stars[5], "4星": stars[4], "3星": stars[3],
            "2星": stars[2], "1星": stars[1], "有图评论": sum(bool(row["image_urls"]) for row in rows),
            "有视频评论": sum(bool(row["video_urls"]) for row in rows), "图片嵌入数": sum(len(row["image_paths"]) for row in rows),
            "状态": "OK" if rows else "无评论结果",
        })

    workbook = Workbook()
    reviews_sheet = workbook.active
    reviews_sheet.title = "all_reviews"
    headers = ["全局序号", "ASIN", "ASIN内序号", "评论ID", "评分", "标题", "评论内容", "评论人", "评论日期", "是否认证购买", "Vine", "Helpful", "图片数量", "图片1", "图片2", "图片3", "图片URL", "视频URL", "来源"]
    reviews_sheet.append(headers)
    add_header_style(reviews_sheet)
    for row in all_rows:
        reviews_sheet.append([
            row["global_index"], row["asin"], row["asin_index"], row["id"], row["rating"], row["title"], row["body"],
            row["author"], row["date"], row["verified"], row["vine"], row["helpful"], len(row["image_urls"]), "", "", "",
            "\n".join(row["image_urls"]), "\n".join(row["video_urls"]), row["source"],
        ])
    widths = {"A":10,"B":14,"C":10,"D":18,"E":8,"F":34,"G":78,"H":18,"I":34,"J":16,"K":10,"L":12,"M":10,"N":18,"O":18,"P":18,"Q":58,"R":46,"S":18}
    for column, width in widths.items():
        reviews_sheet.column_dimensions[column].width = width
    for excel_row, row in enumerate(all_rows, 2):
        reviews_sheet.row_dimensions[excel_row].height = 76 if row["image_paths"] else 46
        for cell in reviews_sheet[excel_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column, path in zip(["N", "O", "P"], row["image_paths"][:3]):
            image = make_excel_image(path)
            if image:
                reviews_sheet.add_image(image, f"{column}{excel_row}")
    reviews_sheet.freeze_panes = "A2"
    reviews_sheet.auto_filter.ref = reviews_sheet.dimensions

    product_csv = input_dir / "product_comparison.csv"
    if product_csv.exists():
        with product_csv.open(encoding="utf-8-sig", newline="") as handle:
            product_rows = list(csv.DictReader(handle))
        if product_rows:
            product_sheet = workbook.create_sheet("product_comparison")
            product_headers = list(product_rows[0].keys())
            product_sheet.append(product_headers)
            add_header_style(product_sheet)
            for row in product_rows:
                product_sheet.append([safe_text(row.get(header, "")) for header in product_headers])
            for index in range(1, len(product_headers) + 1):
                product_sheet.column_dimensions[product_sheet.cell(1, index).column_letter].width = 20
            for row in product_sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            product_sheet.freeze_panes = "A2"
            product_sheet.auto_filter.ref = product_sheet.dimensions

    summary_sheet = workbook.create_sheet("reviews_summary")
    summary_headers = list(summary[0].keys()) if summary else ["ASIN", "评论数", "状态"]
    summary_sheet.append(summary_headers)
    add_header_style(summary_sheet)
    for row in summary:
        summary_sheet.append([row.get(header, "") for header in summary_headers])
    for index in range(1, len(summary_headers) + 1):
        summary_sheet.column_dimensions[summary_sheet.cell(1, index).column_letter].width = 16
    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions

    readme = workbook.create_sheet("README")
    readme.append(["项目", "说明"])
    add_header_style(readme)
    readme.append(["ASIN 数", len(asin_dirs)])
    readme.append(["Written reviews", len(all_rows)])
    readme.append(["说明", "评论数是 written reviews，不等于 Amazon 页面显示的总 ratings。"])
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 100

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    (input_dir / "reviews_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "asins": len(asin_dirs), "reviews": len(all_rows), "summary": summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
