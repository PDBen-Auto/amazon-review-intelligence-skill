#!/usr/bin/env python3
import argparse
import json
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("Install openpyxl before validation: pip install openpyxl") from exc


def header_map(sheet):
    return {str(cell.value): index for index, cell in enumerate(sheet[1], 1)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate an Amazon review delivery workbook")
    parser.add_argument("workbook")
    args = parser.parse_args()
    path = Path(args.workbook).resolve()
    issues = []
    workbook = load_workbook(path, read_only=False)

    for required in ["all_reviews", "reviews_summary"]:
        if required not in workbook.sheetnames:
            issues.append(f"Missing sheet: {required}")
    if issues:
        print(json.dumps({"ok": False, "issues": issues}, ensure_ascii=False, indent=2))
        raise SystemExit(1)

    reviews = workbook["all_reviews"]
    summary = workbook["reviews_summary"]
    review_headers = header_map(reviews)
    summary_headers = header_map(summary)
    required_review_headers = ["ASIN", "评论ID", "评分", "评论内容", "图片URL", "视频URL"]
    for header in required_review_headers:
        if header not in review_headers:
            issues.append(f"Missing all_reviews header: {header}")

    suspicious = []
    for sheet in [reviews, summary]:
        for cell in sheet[1]:
            value = str(cell.value or "")
            if "\ufffd" in value or "??" in value:
                suspicious.append(f"{sheet.title}:{value}")
    if suspicious:
        issues.append("Suspicious header encoding: " + ", ".join(suspicious))

    review_count = max(reviews.max_row - 1, 0)
    summary_total = 0
    if "评论数" in summary_headers:
        column = summary_headers["评论数"]
        for row in range(2, summary.max_row + 1):
            summary_total += int(summary.cell(row, column).value or 0)
        if review_count != summary_total:
            issues.append(f"Row mismatch: all_reviews={review_count}, summary={summary_total}")

    duplicate_ids = []
    if "ASIN" in review_headers and "评论ID" in review_headers:
        seen = Counter()
        for row in range(2, reviews.max_row + 1):
            asin = str(reviews.cell(row, review_headers["ASIN"]).value or "")
            review_id = str(reviews.cell(row, review_headers["评论ID"]).value or "")
            if review_id:
                seen[(asin, review_id)] += 1
        duplicate_ids = [f"{asin}:{review_id}" for (asin, review_id), count in seen.items() if count > 1]
        if duplicate_ids:
            issues.append(f"Duplicate review IDs: {len(duplicate_ids)}")

    result = {
        "ok": not issues,
        "file": str(path),
        "sheets": workbook.sheetnames,
        "review_rows": review_count,
        "summary_total": summary_total,
        "embedded_review_images": len(reviews._images),
        "issues": issues,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if issues:
        sys.exit(1)


if __name__ == "__main__":
    main()
