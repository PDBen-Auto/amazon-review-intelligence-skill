"""Export enriched review JSON to a single-sheet XLSX. Requires openpyxl."""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ['序号','评论人','星级','主类归因','日期','标题','原文','中文摘要','变体/规格','Verified','Helpful','来源','原评论链接']

def export(src, dst):
    data = json.loads(Path(src).read_text(encoding='utf-8-sig'))
    rows = data['reviews']
    market = data.get('product', {}).get('marketplace', 'US')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Review原始数据'
    ws.append(HEADERS)
    for i, r in enumerate(rows, 1):
        rating = r.get('rating')
        if rating is not None and (type(rating) is not int or rating not in range(1, 6)):
            raise ValueError(f'Review {i}: invalid rating')
        verified = r.get('verified_purchase')
        if verified is not None and type(verified) is not bool:
            raise ValueError(f'Review {i}: verified_purchase must be bool/null')
        values = [i,r.get('author'),rating,r.get('primary_issue'),r.get('date') or r.get('date_raw'),r.get('title'),r.get('text'),r.get('summary_zh'),r.get('variant'),None if verified is None else ('Yes' if verified else 'No'),r.get('helpful_votes'),f'Amazon {market} Review',r.get('review_url')]
        for j, value in enumerate(values, 1):
            if value is not None and not isinstance(value, (str, int, float, bool)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, str) and len(value) > 32767:
                raise ValueError(f'Review {i}, column {j}: exceeds Excel cell limit; not truncated')
            cell = ws.cell(i+1,j,value)
            if isinstance(value,str):
                cell.data_type = 's'  # untrusted review text must never become a formula
            cell.alignment = Alignment(vertical='top',wrap_text=True)
    for c in ws[1]:
        c.font = Font(color='FFFFFF',bold=True)
        c.fill = PatternFill('solid',fgColor='17365D')
    for j,width in enumerate([8,20,8,24,24,36,80,55,25,12,10,23,48],1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    Path(dst).parent.mkdir(parents=True,exist_ok=True)
    wb.save(dst)
    check = load_workbook(dst)
    assert check.sheetnames == ['Review原始数据']
    assert check.active.max_row == len(rows)+1
    for i,r in enumerate(rows,2):
        assert (check.active.cell(i,7).value or '') == (r.get('text') or '')
        assert check.active.cell(i,7).data_type != 'f'
    print(f'Exported and verified {len(rows)} reviews: {dst}')

if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--input',required=True)
    p.add_argument('--output',required=True)
    args = p.parse_args()
    export(args.input,args.output)
