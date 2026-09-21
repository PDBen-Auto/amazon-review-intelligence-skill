import json
import re
import subprocess
import sys
import tempfile
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "examples" / "synthetic-demo" / "reviews.json"
XLSX_PATH = ROOT / "examples" / "synthetic-demo" / "amazon-review-demo.xlsx"
HTML_PATH = ROOT / "examples" / "synthetic-review-insight.html"


class DemoParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids = []
        self.scripts = []
        self._script = None

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if "id" in attributes:
            self.ids.append(attributes["id"])
        if tag == "script" and "src" not in attributes and "type" not in attributes:
            self._script = []

    def handle_data(self, data):
        if self._script is not None:
            self._script.append(data)

    def handle_endtag(self, tag):
        if tag == "script" and self._script is not None:
            self.scripts.append("".join(self._script))
            self._script = None


def main() -> None:
    subprocess.run([sys.executable, str(ROOT / "examples" / "synthetic-demo" / "build_demo.py")], check=True)
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    assert data["collection"]["synthetic"] is True
    assert len(data["reviews"]) == 20
    assert len({review["review_id"] for review in data["reviews"]}) == 20

    workbook = load_workbook(XLSX_PATH, read_only=False)
    sheet = workbook["Review原始数据"]
    assert sheet.max_row == 21
    assert sheet.cell(10, 4).value == "Strong airflow"
    assert sheet.cell(10, 7).data_type != "f"

    markup = HTML_PATH.read_text(encoding="utf-8")
    parser = DemoParser()
    parser.feed(markup)
    assert len(parser.ids) == len(set(parser.ids)), "HTML ids must be unique"
    assert parser.scripts, "Interactive demo script is missing"
    assert "SYNTHETIC DATA" in markup
    assert "const APP =" in parser.scripts[-1]
    assert re.search(r"<meta name=\"description\"", markup)

    with tempfile.NamedTemporaryFile("w", suffix=".js", encoding="utf-8", delete=False) as handle:
        handle.write(parser.scripts[-1])
        script_path = Path(handle.name)
    try:
        subprocess.run(["node", "--check", str(script_path)], check=True)
    finally:
        script_path.unlink(missing_ok=True)
    print("synthetic demo JSON, XLSX, HTML, and JavaScript: OK")


if __name__ == "__main__":
    main()
