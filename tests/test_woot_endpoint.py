import importlib.util
import json
import tempfile
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("amazon_review_scraper", ROOT / "scripts" / "amazon_review_scraper.py")
scraper = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(scraper)

ASIN = "B07TPPVYLQ"
OFFER = "7eca00f4-6e22-46a0-8ee4-1a581bfacb75"
URL = "https://home.woot.com/offers/test-product"
PAGE = f'<script>var offerItems = [{{"Id":"item-1","Asin":"{ASIN}","Attributes":[{{"Key":"Color","Value":"Black"}}]}}]; var offerId = "{OFFER}";</script>'


def main():
    with patch.object(scraper, "fetch_text", return_value=PAGE):
        source = scraper.resolve_woot_offer(ASIN, URL, 10, 0)
        assert source["offer_id"] == OFFER
        assert source["review_base"] == f"https://home.woot.com/review/Reviews/{ASIN}"
        try:
            scraper.resolve_woot_offer("B0CM6S7VT5", URL, 10, 0)
            raise AssertionError("mismatched ASIN accepted")
        except ValueError as exc:
            assert "ASIN mismatch" in str(exc)

    called = []
    full_body = "A long review. " * 200

    def response(url, timeout, retries):
        called.append(url)
        if "pagingNext=" in url:
            return json.dumps({"Reviews": [{"Title": "Second page", "Text": "Second body", "OverallRating": 4}], "PagingNext": ""})
        return json.dumps({"Reviews": [{"Title": "First page", "Text": full_body, "OverallRating": 5}], "PagingNext": "cursor-two"})

    with patch.object(scraper, "fetch_text", side_effect=response):
        rows = scraper.fetch_reviews(source, timeout=10, retries=0)
    assert len(rows) == 2
    assert rows[0]["Text"] == full_body
    assert f"offerid={OFFER}" in called[0]
    assert called[0].startswith(source["review_base"])
    assert "pagingNext=cursor-two" in called[1]

    def interrupted(url, timeout, retries):
        if "pagingNext=" in url:
            raise TimeoutError("read timed out")
        return json.dumps({"Reviews": [{"Title": "First page", "Text": full_body}], "PagingNext": "cursor-two"})

    errors = []
    with patch.object(scraper, "fetch_text", side_effect=interrupted):
        partial = scraper.fetch_reviews(source, timeout=10, retries=0, errors=errors)
    assert len(partial) == 1
    assert errors[0]["page"] == 2 and "timed out" in errors[0]["error"]

    with patch.object(scraper, "fetch_text", side_effect=interrupted):
        captured, scrape_errors = scraper.scrape(source, "basic", 0, 10, 0)
    assert len(captured) == 1 and scrape_errors[0]["page"] == 2

    with patch.object(scraper, "fetch_text", return_value="<html>404</html>"):
        try:
            scraper.fetch_reviews(source, timeout=10, retries=0)
            raise AssertionError("invalid JSON accepted")
        except json.JSONDecodeError:
            pass

    with tempfile.TemporaryDirectory() as tmp:
        from subprocess import run
        import sys
        mapping = Path(tmp) / "map.json"
        mapping.write_text("{}", encoding="utf-8")
        asin_dir = Path(tmp) / ASIN
        asin_dir.mkdir()
        (asin_dir / f"{ASIN}_woot_max.json").write_text(json.dumps({"summary": {"status": "collected"}, "reviews": [{"Text": "old"}]}), encoding="utf-8")
        result = run([sys.executable, str(ROOT / "scripts" / "batch_review_collect.py"), ASIN,
                      "--woot-offer-map", str(mapping), "--output-dir", tmp, "--mode", "max", "--resume"], capture_output=True, text=True)
        assert result.returncode == 0
        status = json.loads((Path(tmp) / "collection_status.json").read_text(encoding="utf-8"))[0]
        assert status["status"] == "woot_mapping_missing"
        assert status["previousCaptureRetained"] is True
        workbook = Path(tmp) / "reviews.xlsx"
        built = run([sys.executable, str(ROOT / "scripts" / "build_review_delivery.py"),
                     "--input-dir", tmp, "--output", str(workbook)], capture_output=True,
                    text=True, encoding="utf-8", errors="replace")
        assert built.returncode == 0, built.stderr
        from openpyxl import load_workbook
        book = load_workbook(workbook, read_only=True)
        sheet = book["reviews_summary"]
        headers = [cell.value for cell in sheet[1]]
        assert "旧采集数据" in sheet.cell(2, headers.index("状态") + 1).value
        book.close()
    print("woot endpoint fixture: OK")


if __name__ == "__main__":
    main()
