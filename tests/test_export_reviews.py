import json
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = Path(__file__).resolve().parent / "fixtures" / "minimal.json"


def main() -> None:
    output = ROOT / "tests" / "_tmp_review_export.xlsx"
    command = [
        sys.executable,
        str(ROOT / "scripts" / "export_reviews.py"),
        "--input",
        str(FIXTURE),
        "--output",
        str(output),
    ]
    subprocess.run(command, check=True)
    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=False)
    assert workbook.sheetnames == ["Review原始数据"]
    sheet = workbook.active
    assert sheet.max_row == 2
    assert sheet.cell(2, 7).value == "The fit is solid and setup was quick."
    assert sheet.cell(2, 7).data_type != "f"
    output.unlink()
    print("export_reviews fixture: OK")


if __name__ == "__main__":
    main()
